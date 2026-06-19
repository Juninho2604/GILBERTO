"""Carga del histórico de refinación (53 lotes) desde el ``.xlsx``.

Cada lote es un bloque de 13 filas en la hoja ``Hoja1``:

    Receptive Information ...                         (encabezado)
    Arrival date     | <fecha>
    Receptive date   | <fecha> | CU | Kg.  | grade | content | RR | ...
    Customer lot No. | <n>     | AU | g/t  | grade | ...
    Lot No. (JX)     | <n>     | AG | g/t  | grade | ...
                     |         | PT | g/t  | grade | ...
    WMT              | <kg>    | PD | g/t  | grade | ...
    Moisture         | <frac>
    DMT              | <kg>
    Resultado USD/Kg | ...
    <receta>                                          (string de mezcla)

Devuelve objetos :class:`HistoricLot` con pesos, leyes reales y receta parseada.
Estos son el insumo de la estimación de leyes (sección 6).
"""

from __future__ import annotations

from dataclasses import dataclass, field
from pathlib import Path
from typing import Optional

from openpyxl import load_workbook

from domain.models import METALS

from .recipes import ParsedRecipe, parse_recipe

DEFAULT_HISTORY_XLSX = (
    Path(__file__).resolve().parents[2] / "data" / "refining_history.xlsx"
)

_METAL_COL = 3   # columna D: nombre del metal
_GRADE_COL = 5   # columna F: ley


@dataclass
class HistoricLot:
    """Un lote histórico enviado a la refinería."""

    customer_lot: object
    jx_lot: object
    wmt: float
    moisture: float
    dmt: float
    grades: dict[str, float]            # ley real del lote por metal
    recipe_raw: str
    recipe: ParsedRecipe = field(repr=False)

    @property
    def index_codes(self) -> list[str]:
        return self.recipe.codes


def _label_map(block) -> dict[str, int]:
    out: dict[str, int] = {}
    for ri, row in enumerate(block):
        if row and row[0] is not None:
            out[str(row[0]).strip()] = ri
    return out


def load_history(path: str | Path = DEFAULT_HISTORY_XLSX) -> list[HistoricLot]:
    wb = load_workbook(path, data_only=True)
    ws = wb["Hoja1"] if "Hoja1" in wb.sheetnames else wb.worksheets[0]
    rows = list(ws.iter_rows(values_only=True))

    starts = [i for i, r in enumerate(rows) if r and r[0] == "Receptive Information"]
    lots: list[HistoricLot] = []
    for s in starts:
        block = rows[s : s + 13]
        labels = _label_map(block)

        def lab(key: str) -> Optional[object]:
            ri = labels.get(key)
            return block[ri][1] if ri is not None else None

        grades: dict[str, float] = {}
        for row in block:
            if row and row[_METAL_COL] in METALS:
                val = row[_GRADE_COL]
                grades[row[_METAL_COL]] = float(val) if isinstance(val, (int, float)) else 0.0

        # La receta es la fila siguiente a "Resultado USD / Kg".
        recipe_raw = ""
        res_ri = labels.get("Resultado USD / Kg")
        if res_ri is not None and res_ri + 1 < len(block):
            cell = block[res_ri + 1][0]
            recipe_raw = str(cell).strip() if cell else ""

        wmt = lab("WMT")
        moisture = lab("Moisture")
        dmt = lab("DMT")
        lots.append(
            HistoricLot(
                customer_lot=lab("Customer lot No."),
                jx_lot=lab("Lot No. (JX)"),
                wmt=float(wmt) if isinstance(wmt, (int, float)) else 0.0,
                moisture=float(moisture) if isinstance(moisture, (int, float)) else 0.0,
                dmt=float(dmt) if isinstance(dmt, (int, float)) else 0.0,
                grades=grades,
                recipe_raw=recipe_raw,
                recipe=parse_recipe(recipe_raw),
            )
        )
    return lots


if __name__ == "__main__":  # pragma: no cover
    lots = load_history()
    print(f"Lotes cargados: {len(lots)}")
    resolvable = sum(1 for L in lots if L.recipe.resolvable)
    print(f"Recetas resolubles: {resolvable} | ambiguas/texto: {len(lots) - resolvable}")
    for L in lots[:8]:
        print(f"  cust={L.customer_lot} WMT={L.wmt:>6.0f}  {L.recipe_raw:<28} {L.recipe.note}")
