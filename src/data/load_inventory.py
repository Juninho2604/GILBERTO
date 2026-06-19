"""Carga del inventario real desde el ``.xlsx`` (sección 9, paso 3 del brief).

El archivo ``inventory.xlsx`` tiene una hoja con bloques por categoría:

    630 - Inventory RAEE
    001  Bajo Grado Marron   Tracked   13329.97
    ...
    Total 630 - Inventory RAEE
    650 - Inventario No Ferroso
    ...

Solo la categoría **630 (RAEE)** entra en las mezclas. Las leyes no están en
este archivo: se estiman (``estimate_grades.py``) o se cargan por laboratorio.
"""

from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook

from domain.models import Category, InventoryItem

# Encabezados de categoría tal como aparecen en la hoja.
_CATEGORY_HEADERS = {
    "630 - Inventory RAEE": Category.RAEE,
    "650 - Inventario No Ferroso": Category.NON_FERROUS,
    "651 - Inventario Ferroso": Category.FERROUS,
}

DEFAULT_INVENTORY_XLSX = Path(__file__).resolve().parents[2] / "data" / "inventory.xlsx"


def normalize_code(raw: object) -> str:
    """Normaliza un código a 3 dígitos (``6`` → ``"006"``) para casar recetas.

    Las recetas históricas referencian códigos sin ceros a la izquierda
    (``6``, ``11``, ``998``); el inventario los guarda con ceros (``006``).
    """
    s = str(raw).strip()
    if s.isdigit():
        return f"{int(s):03d}"
    return s


def load_inventory(path: str | Path = DEFAULT_INVENTORY_XLSX) -> list[InventoryItem]:
    """Devuelve todos los ítems del inventario, con su categoría.

    Las leyes quedan en 0 (se completan luego). La humedad usa el default del
    modelo. ``quantity_kg`` es el stock disponible (WMT).
    """
    wb = load_workbook(path, data_only=True)
    ws = wb.worksheets[0]

    items: list[InventoryItem] = []
    current: Category | None = None
    for row in ws.iter_rows(values_only=True):
        if not row or row[0] is None:
            continue
        first = str(row[0]).strip()

        if first in _CATEGORY_HEADERS:
            current = _CATEGORY_HEADERS[first]
            continue
        if first.startswith("Total") or first in ("Item Code",) or current is None:
            continue
        # Encabezados generales (título, empresa, fecha) → sin nombre/cantidad.
        name = row[1] if len(row) > 1 else None
        qty = row[3] if len(row) > 3 else None
        if name is None or not isinstance(qty, (int, float)):
            continue

        items.append(
            InventoryItem(
                code=normalize_code(first),
                name=str(name).strip(),
                category=current,
                quantity_kg=float(qty),
            )
        )
    return items


def load_raee_inventory(
    path: str | Path = DEFAULT_INVENTORY_XLSX, with_stock_only: bool = False
) -> list[InventoryItem]:
    """Solo los ítems RAEE (alcance del simulador). Opcionalmente con stock > 0."""
    items = [it for it in load_inventory(path) if it.category == Category.RAEE]
    if with_stock_only:
        items = [it for it in items if it.quantity_kg > 0]
    return items


if __name__ == "__main__":  # pragma: no cover
    items = load_inventory()
    raee = [i for i in items if i.category == Category.RAEE]
    with_stock = [i for i in raee if i.quantity_kg > 0]
    print(f"Total ítems: {len(items)} | RAEE: {len(raee)} | RAEE con stock: {len(with_stock)}")
    for it in sorted(with_stock, key=lambda i: -i.quantity_kg)[:10]:
        print(f"  {it.code}  {it.name:<45} {it.quantity_kg:>10,.1f} kg")
